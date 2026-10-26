#!/usr/bin/env python3
"""按《投行 PPT 风格规范》(ib-style-guide.md) 生成空白模板 deck。

用法:
    pip install python-pptx
    python scripts/build_template.py [输出路径]

默认输出 templates/ib-template.pptx。模板共 7 页:
封面 / 保密声明 / 章节页 / 内容页(结论式标题) / 表格页 / 图表页 / 结尾页。
规范要点: 16:9、深蓝 #0B2C5C 主色、Arial + 微软雅黑、结论式标题、
每页 Source 脚注 + Confidential 页脚 + 页码。
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---- 规范色值（ib-style-guide.md §4，全稿锁定）----
NAVY = RGBColor(0x0B, 0x2C, 0x5C)
AUX_BLUE = RGBColor(0x1F, 0x4E, 0x79)
LIGHT_BLUE = RGBColor(0xD9, 0xE2, 0xEC)
INK = RGBColor(0x14, 0x24, 0x3D)
MUTED = RGBColor(0x5A, 0x6B, 0x84)
HAIRLINE = RGBColor(0xC9, 0xD2, 0xDE)
ZEBRA = RGBColor(0xF2, 0xF4, 0xF8)
ACCENT = RGBColor(0xA6, 0x19, 0x2E)
GOLD = RGBColor(0xB9, 0x97, 0x5B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FOOT_GRAY = RGBColor(0x6E, 0x7B, 0x8D)

LATIN_FONT = "Arial"
EA_FONT = "Microsoft YaHei"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.67)          # §2 页边距
CONTENT_W = Inches(12.0)       # 正文区宽（13.333 - 2*0.67 ≈ 12.0）

CONFIDENTIAL = "Confidential — For discussion purposes only"


def set_font(run, size, bold=False, color=INK):
    """同时设置西文字体与东亚字体（规范 §3 只用一个字体家族）。"""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = LATIN_FONT
    rPr = run._r.get_or_add_rPr()
    for tag, typeface in (("a:latin", LATIN_FONT), ("a:ea", EA_FONT), ("a:cs", LATIN_FONT)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", typeface)


def add_text(slide, x, y, w, h, text, size, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.45):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        set_font(run, size, bold, color)
    return box


def add_rect(slide, x, y, w, h, color, line_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def footer(slide, page_no, source="Source: 示例数据（演示口径）"):
    """§8 页脚: 左 = 来源, 中 = 保密字样, 右 = 页码; 位置逐页固定。"""
    add_text(slide, MARGIN, Inches(6.92), Inches(6.4), Inches(0.3),
             source, 8.25, color=FOOT_GRAY)
    add_text(slide, Inches(4.87), Inches(6.92), Inches(3.6), Inches(0.3),
             CONFIDENTIAL, 8.25, color=FOOT_GRAY, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(12.33), Inches(6.92), Inches(0.33), Inches(0.3),
             str(page_no), 8.25, color=FOOT_GRAY, align=PP_ALIGN.RIGHT)


def header(slide, title, kicker=None):
    """内容页页眉: kicker(可选) + 结论式标题 + 3px 主色分隔线（§2/§6.1）。"""
    y = Inches(0.38)
    if kicker:
        add_text(slide, MARGIN, y, CONTENT_W, Inches(0.28),
                 kicker, 10.5, bold=True, color=MUTED)
        y = Inches(0.64)
    add_text(slide, MARGIN, y, CONTENT_W, Inches(0.75),
             title, 21, bold=True, color=NAVY)
    add_rect(slide, MARGIN, Inches(1.28), CONTENT_W, Pt(3), NAVY)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_cover(prs):
    s = blank(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    add_rect(s, MARGIN, Inches(2.02), Inches(1.6), Pt(2.25), GOLD)
    add_text(s, MARGIN, Inches(2.25), Inches(11.0), Inches(1.2),
             "项目名称 — 交易概览", 40, bold=True, color=WHITE)
    add_text(s, MARGIN, Inches(3.35), Inches(11.0), Inches(0.6),
             "Project Name — Transaction Overview", 15, color=LIGHT_BLUE)
    add_text(s, MARGIN, Inches(6.55), Inches(8.0), Inches(0.4),
             "呈送：董事会 / 投资委员会", 12, color=LIGHT_BLUE)
    add_text(s, Inches(9.83), Inches(6.55), Inches(2.83), Inches(0.4),
             "2026 年 10 月", 12, color=LIGHT_BLUE, align=PP_ALIGN.RIGHT)
    add_text(s, MARGIN, Inches(6.98), Inches(11.99), Inches(0.3),
             CONFIDENTIAL, 8.25, color=LIGHT_BLUE)
    return s


def slide_disclaimer(prs):
    s = blank(prs)
    header(s, "保密声明与重要提示")
    body = (
        "本文件仅供讨论使用，所载信息基于公开资料与管理层提供的数据整理，未经独立审计。\n"
        "本文件包含的预测与展望性陈述（以 E 标注）基于当前假设，实际结果可能存在重大差异。\n"
        "未经许可，不得复制、转发或向任何第三方披露本文件内容。"
    )
    add_text(s, MARGIN, Inches(1.7), CONTENT_W, Inches(3.5), body, 13.5,
             color=INK, line_spacing=1.55)
    footer(s, 2, source="Source: —")
    return s


def slide_section(prs, number="01", title="行业与市场分析", page_no=3):
    s = blank(prs)
    add_text(s, MARGIN, Inches(2.6), Inches(2.2), Inches(1.1),
             number, 54, bold=True, color=LIGHT_BLUE)
    add_text(s, Inches(2.9), Inches(2.78), Inches(9.0), Inches(0.8),
             title, 28, bold=True, color=NAVY)
    add_rect(s, Inches(2.93), Inches(3.62), Inches(1.6), Pt(2.25), GOLD)
    footer(s, page_no, source="Source: —")
    return s


def slide_content(prs):
    s = blank(prs)
    header(s, "营收连续四季增长，四季度达 61 亿元（结论式标题）",
           kicker="财务分析 · FINANCIAL ANALYSIS")
    bullets = [
        ("增长来源", "增量主要来自替换胎市场放量与高毛利产品占比提升，单季毛利率环比 +1.8pct。"),
        ("口径提示", "本页数字为示例数据，仅用于演示版式；正式稿须逐项标注来源与期间口径。"),
        ("写作规则", "每页只讲一个结论；标题写完整结论句，正文只做证据与标注，不写三段以上文字。"),
    ]
    y = 1.7
    for lead, rest in bullets:
        box = s.shapes.add_textbox(MARGIN, Inches(y), CONTENT_W, Inches(0.9))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.line_spacing = 1.5
        r0 = p.add_run(); r0.text = "▪ "  # 项目符号
        set_font(r0, 13.5, bold=True, color=NAVY)
        r1 = p.add_run(); r1.text = lead + "："
        set_font(r1, 13.5, bold=True, color=INK)
        r2 = p.add_run(); r2.text = rest
        set_font(r2, 13.5, color=INK)
        y += 1.05
    # KPI 卡片行（隔行底色块，无阴影无圆角堆叠）
    kpis = [("61 亿元", "四季度营收（示例）"), ("+12.4%", "同比增速（示例）"), ("18.6%", "毛利率（示例）")]
    x = 0.67
    for value, label in kpis:
        add_rect(s, Inches(x), Inches(5.1), Inches(3.84), Inches(1.3), ZEBRA)
        add_text(s, Inches(x + 0.3), Inches(5.32), Inches(3.3), Inches(0.55),
                 value, 20, bold=True, color=NAVY)
        add_text(s, Inches(x + 0.3), Inches(5.92), Inches(3.3), Inches(0.35),
                 label, 10.5, color=MUTED)
        x += 4.08
    footer(s, 4)
    return s


def slide_table(prs):
    s = blank(prs)
    header(s, "可比公司：头部企业收入规模与盈利能力对比（示例数据）",
           kicker="可比公司分析 · COMPARABLES")
    rows, cols = 5, 4
    gt = s.shapes.add_table(rows, cols, MARGIN, Inches(1.7),
                            CONTENT_W, Inches(3.4))
    table = gt.table
    widths = [Inches(3.6), Inches(2.8), Inches(2.8), Inches(2.8)]
    for i, w in enumerate(widths):
        table.columns[i].width = w
    data = [
        ("公司", "营收（亿元）", "同比增速", "毛利率"),
        ("公司 A", "245.6", "+9.8%", "21.3%"),
        ("公司 B", "187.2", "+6.1%", "19.8%"),
        ("公司 C", "132.9", "+12.4%", "18.6%"),
        ("行业平均", "188.6", "+9.4%", "19.9%"),
    ]
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
            elif r == rows - 1:
                cell.fill.fore_color.rgb = WHITE
            else:
                cell.fill.fore_color.rgb = ZEBRA if r % 2 == 0 else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = val
            bold = r == 0 or r == rows - 1
            color = WHITE if r == 0 else INK
            set_font(run, 12, bold=bold, color=color)
    add_text(s, MARGIN, Inches(5.35), CONTENT_W, Inches(0.4),
             "注：单位写进表头，数字右对齐、千分位、小数位全表统一；合计/均值行上方加主色线、数字加粗（§5）。",
             10.5, color=MUTED)
    footer(s, 5)
    return s


def slide_chart(prs):
    s = blank(prs)
    header(s, "季度营收稳步爬坡，四季度创单季新高（示例数据）",
           kicker="财务分析 · FINANCIAL ANALYSIS")
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("营收（亿元）", (48.2, 52.7, 57.3, 61.0))
    gf = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.65),
        Inches(8.6), Inches(4.6), chart_data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    series = chart.plots[0].series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = NAVY
    chart.plots[0].has_data_labels = True
    dl = chart.plots[0].data_labels
    dl.show_value = True
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10.5)
    dl.font.name = LATIN_FONT
    dl.font.color.rgb = MUTED
    va = chart.value_axis
    va.minimum_scale = 0  # §6.3 柱图从零基线起
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = HAIRLINE
    va.tick_labels.font.size = Pt(10.5)
    va.tick_labels.font.name = LATIN_FONT
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(10.5)
    ca.tick_labels.font.name = LATIN_FONT
    # 右侧要点栏
    add_rect(s, Inches(9.6), Inches(1.65), Inches(0.03), Inches(4.6), HAIRLINE)
    add_text(s, Inches(9.85), Inches(1.7), Inches(2.8), Inches(4.4),
             "要点\n\n· 同系深浅配色，焦点柱才用强调红\n· 关键点必标数据标签\n· 截断坐标轴必须显式标注断轴",
             12, color=INK, line_spacing=1.5)
    footer(s, 6)
    return s


def slide_closing(prs):
    s = blank(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    add_rect(s, MARGIN, Inches(2.9), Inches(1.6), Pt(2.25), GOLD)
    add_text(s, MARGIN, Inches(3.15), Inches(11.0), Inches(1.0),
             "下一步与时间表", 32, bold=True, color=WHITE)
    add_text(s, MARGIN, Inches(4.05), Inches(11.0), Inches(0.6),
             "Next steps — 管理层确认估值区间后，两周内提交正式 pitch book。",
             13.5, color=LIGHT_BLUE)
    add_text(s, MARGIN, Inches(6.98), Inches(11.99), Inches(0.3),
             CONFIDENTIAL, 8.25, color=LIGHT_BLUE)
    add_text(s, Inches(12.33), Inches(6.98), Inches(0.33), Inches(0.3),
             "7", 8.25, color=LIGHT_BLUE, align=PP_ALIGN.RIGHT)
    return s


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else \
        Path(__file__).resolve().parent.parent / "templates" / "ib-template.pptx"
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_cover(prs)
    slide_disclaimer(prs)
    slide_section(prs)
    slide_content(prs)
    slide_table(prs)
    slide_chart(prs)
    slide_closing(prs)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out} ({len(prs.slides.slides if hasattr(prs.slides,'slides') else prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
